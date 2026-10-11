import re,json,argparse,subprocess,shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
parser=argparse.ArgumentParser(description='Generar el manual Word y, opcionalmente, PDF')
parser.add_argument('--output-dir',default='manual-exportado')
parser.add_argument('--pdf',action='store_true',help='Requiere LibreOffice y pypdf')
args=parser.parse_args()
root=Path(__file__).resolve().parents[1]
source=(root/'docs/manual-maderas-del-mundo.md').read_text()
out=Path(args.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
page_map={}
version=re.search(r'Versión (\d+)',source).group(1)
date=re.search(r'Versión \d+ · (.+)',source).group(1)
filename=f'Manual_Maderas_del_Mundo_V{version}'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.line_spacing=1.06
normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.widow_control=True
for name,size in [('Title',27),('Subtitle',12),('Heading 1',17),('Heading 2',12.5),('Heading 3',11)]:
 s=doc.styles[name];s.font.size=Pt(size)
 if name.startswith('Heading'):
  s.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 9)
  s.paragraph_format.keep_with_next=True
# footer required for long reference manual
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run(f'Maderas del Mundo · Versión {version} · ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(80,80,80)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)

def text_runs(p,txt):
 # plain readable inline text; hyperlinks handled for bare URLs
 for part in re.split(r'(https?://[^\s]+)',txt.replace('`','')):
  if part.startswith('http'):
   rel=doc.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1A4F73');pr.append(c);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=part;r.append(t);h.append(r);p._p.append(h)
  else:p.add_run(part)

def bookmark(p,name,num):
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(num));start.set(qn('w:name'),name)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(num));p._p.insert(0,start);p._p.append(end)

def toc_link(label,num):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),f'cap{num}')
 r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label + (f'    {page_map[str(num)]}' if str(num) in page_map else '');r.append(t);h.append(r);p._p.append(h)

lines=source.splitlines();doc.add_paragraph('Maderas del Mundo',style='Title');doc.add_paragraph('Manual de administración y mantenimiento',style='Subtitle')
for text in [f'Versión {version}',date,'Luis Rodríguez']:
 p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(5)
doc.add_paragraph('Una guía para modificar el catálogo, las fotografías, el inventario y la venta, y mantener conectadas todas las plataformas.')
doc.add_paragraph('Web pública');text_runs(doc.add_paragraph(),'https://www.maderasdelmundo.es')
doc.add_paragraph('Código de referencia');text_runs(doc.add_paragraph(),'https://github.com/luisinko-lasko/maderas-del-mundo')
doc.add_paragraph(next(l for l in source.splitlines() if l.startswith('Estado de referencia del código:')))
doc.add_paragraph('Las claves reales se conservan en los paneles y en un gestor de contraseñas.')
doc.add_page_break();doc.add_heading('Índice',level=1)
chapters=[]
for line in lines:
 if line.startswith('## '):
  label=line[3:];num=int(label.split()[0]);chapters.append((num,label));toc_link(label,num)
doc.add_paragraph('Pulsa un capítulo para abrirlo. El capítulo 2 localiza las tareas más frecuentes; el capítulo 14 reúne las soluciones a errores habituales.')
doc.add_page_break();doc.add_heading('Cómo utilizar este manual',level=1)
# opening two paragraphs
intro_start=next(i for i,l in enumerate(lines) if l.startswith('Este manual permite'))
intro_end=next(i for i,l in enumerate(lines) if l.startswith('## 1 '))
for l in lines[intro_start:intro_end]:
 if l.strip():text_runs(doc.add_paragraph(),l)
i=intro_end; in_code=False
while i<len(lines):
 l=lines[i]
 if l.startswith('```'):
  if not in_code:
   code=[];i+=1
   while i<len(lines) and not lines[i].startswith('```'):
    code.append(lines[i]);i+=1
   for c in code:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
    p.paragraph_format.left_indent=Inches(.12)
    r=p.add_run(c if c else ' ');r.font.name='Consolas';r.font.size=Pt(8.3)
   doc.add_paragraph().paragraph_format.space_after=Pt(1)
 elif l.startswith('## '):
  txt=l[3:];p=doc.add_heading(txt,level=1);num=int(txt.split()[0]);bookmark(p,f'cap{num}',num)
 elif l.startswith('### '):doc.add_heading(l[4:],level=2)
 elif l.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in row):rows.append(row)
   i+=1
  i-=1
  cols=len(rows[0]);table=doc.add_table(rows=0,cols=cols);table.autofit=False
  widths=([1.75,2.65,2.6] if cols==3 else [2.5,4.5] if cols==2 else [2.7,.9,.85,.85,.85,.85] if cols==6 else [7/cols]*cols)
  for c,w in zip(table.columns,widths):c.width=Inches(w)
  for ri,row in enumerate(rows):
   cells=table.add_row().cells
   for ci,(cell,txt) in enumerate(zip(cells,row)):
    cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
    tcpr.append(b);m=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');m.append(el)
    tcpr.append(m)
    if ri==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF0');tcpr.append(shade)
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
    text_runs(p,txt)
    for run in p.runs:run.font.size=Pt(9);run.bold=(ri==0)
   pr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
   if ri==0:
    repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(2)
 elif re.match(r'^\d+\. ',l):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);text_runs(p,l)
 elif l.startswith('- '):
  p=doc.add_paragraph(style='List Bullet');text_runs(p,l[2:])
 elif l.strip():text_runs(doc.add_paragraph(),l)
 i+=1
# Remove inherited template borders from all styles and paragraphs
for style in doc.styles:
 for border in list(style.element.findall('.//' + qn('w:pBdr'))):
  border.getparent().remove(border)
for paragraph in doc.paragraphs:
 for border in list(paragraph._p.findall('.//' + qn('w:pBdr'))):
  border.getparent().remove(border)
# keep title page all black, no paragraph borders
for p in doc.paragraphs:
 if p.style.name in ['Title','Subtitle']:
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
doc.core_properties.title='Manual de administración y mantenimiento de Maderas del Mundo'
doc.core_properties.subject=f'Versión {version}'
doc.core_properties.author='Maderas del Mundo'
doc.core_properties.keywords='manual, administración, GitHub, Vercel, Supabase, Stripe, Brevo, MDM'
target=out/(filename+'.docx')
doc.save(target)
if args.pdf:
 lo=shutil.which('libreoffice') or shutil.which('soffice')
 mac='/Applications/LibreOffice.app/Contents/MacOS/soffice'
 if not lo and Path(mac).exists():lo=mac
 if not lo:raise SystemExit('Instala LibreOffice para generar el PDF; el Word ya está creado.')
 def convert():
  subprocess.run([lo,'--headless','--convert-to','pdf','--outdir',str(out),str(target)],check=True)
 convert()
 from pypdf import PdfReader
 pdf=out/(filename+'.pdf')
 pages=PdfReader(pdf).pages
 for number,label in chapters:
  normalized=' '.join(label.split())
  for index,page in enumerate(pages[2:],start=3):
   if normalized in ' '.join((page.extract_text() or '').split()):
    page_map[str(number)]=index;break
 # Add page numbers to the navigation paragraph without changing its line count.
 for p in doc.paragraphs:
  for hyperlink in p._p.findall(qn('w:hyperlink')):
   anchor=hyperlink.get(qn('w:anchor'),'')
   if anchor.startswith('cap'):
    key=anchor[3:]
    if key in page_map:
     t=hyperlink.find('.//'+qn('w:t'));t.text += '    '+str(page_map[key])
 doc.save(target);convert()
 print('PDF:',pdf)
print('Word:',target)
